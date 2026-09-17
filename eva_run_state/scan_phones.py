import os, re, json, time, sys
from concurrent.futures import ThreadPoolExecutor
import pymupdf
import openpyxl

B_TABLE = r'S:\2 人事\4. 简历库\人才登记表20260916-B.xlsx'
SORTED_DIR = r'S:\2 人事\4. 简历库\Sorted'
UNSORTED_DIR = r'S:\2 人事\4. 简历库\Unsorted'
OUTPUT = r'C:\Users\vecto\Documents\4.VTribe\eva_run_state\phone_matches.json'

WORKERS = 20

PATTERNS = [
    re.compile(r'【(?P<title>[^_【】]+?)\s*[_\s](?P<city>[^_【】]*?)\s+(?P<salary>[^】]*?)】\s*(?P<name>[^】\d]+?)\s+(?P<exp>\d+年(?:以内|应届生|以上|毕业)?)(?:\s*-\s*Copy)?\.pdf', re.UNICODE),
    re.compile(r'【(?P<title>[^_【】]+?)\s*[_\s](?P<city>[^_【】]*?)】\s*(?P<name>[^】\d]+?)\s+(?P<exp>\d+年(?:以内|应届生|以上|毕业)?)(?:\s*-\s*Copy)?\.pdf', re.UNICODE),
    re.compile(r'(?P<name>[^_]+)_(?P<age>\d+岁|[\d\.]+)_(?P<title>[^_]+)_(?P<city>[^_]+)_.*\.pdf', re.UNICODE),
    re.compile(r'】\s*(?P<name>[^\d】\n]+?)\s+(?:\d+年|应届|\d+岁)(?:.*?)(?:\s*-\s*Copy)?\.pdf', re.UNICODE),
]

PHONE_RE = re.compile(r'(\(\+?\d+\))?\s*\d{3}[\s\-]?\d{4}[\s\-]?\d{4}')
CHINA_PHONE = re.compile(r'(?:\(\+86\))?\s*1[3-9]\d{9}')
INTL_PHONE = re.compile(r'(?:\(\+\d+\))[\d\s\-]{6,15}')

BAD_NAMES = {'未知', 'N/A', '', 'None', '费女士', '先生', '女士', '糖 糖'}

def extract_name_from_filename(filename):
    for pat in PATTERNS:
        m = pat.search(filename)
        if m:
            name = m.group('name').strip().strip('_- ')
            if name and name not in BAD_NAMES:
                return name
    return None

def normalize_phone(raw):
    digits = re.sub(r'\D', '', raw)
    if len(digits) == 11 and digits.startswith('1'):
        return f'(+86){digits}'
    if len(digits) > 11:
        return f'(+86){digits[-11:]}'
    return None

def load_unknown_phones():
    wb = openpyxl.load_workbook(B_TABLE, read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    name_idx = headers.index('候选人')
    phone_idx = headers.index('电话1')
    
    unknowns = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = str(row[name_idx] or '').strip()
        if name in ('未知', 'N/A', '', 'None'):
            phone = str(row[phone_idx] or '').strip()
            if phone and phone not in ('N/A', '', 'None'):
                unknowns[phone] = True
    
    wb.close()
    print(f'Loaded {len(unknowns)} unknown phone numbers')
    return set(unknowns.keys())

def scan_pdf(args):
    filepath, filename = args
    try:
        doc = pymupdf.open(filepath)
        text = ''
        for page in doc:
            text += page.get_text()
        doc.close()
        
        phones_found = set()
        for m in CHINA_PHONE.finditer(text):
            norm = normalize_phone(m.group(0))
            if norm:
                phones_found.add(norm)
        for m in INTL_PHONE.finditer(text):
            raw = m.group(0).strip()
            digits = re.sub(r'\D', '', raw)
            if len(digits) >= 8:
                phones_found.add(raw)
        
        name_from_filename = extract_name_from_filename(filename)
        
        return {
            'filename': filename,
            'phones': list(phones_found),
            'name_from_filename': name_from_filename,
            'text_preview': text[:100],
        }
    except Exception as e:
        return {
            'filename': filename,
            'phones': [],
            'name_from_filename': None,
            'error': str(e)[:100],
        }

def main():
    start = time.time()
    
    unknown_phones = load_unknown_phones()
    
    files_to_scan = []
    for d, label in [(SORTED_DIR, 'Sorted'), (UNSORTED_DIR, 'Unsorted')]:
        if os.path.isdir(d):
            for f in os.listdir(d):
                if f.lower().endswith('.pdf'):
                    files_to_scan.append((os.path.join(d, f), f))
    
    print(f'Scanning {len(files_to_scan)} PDFs with {WORKERS} threads...')
    
    matches = {}
    scanned = 0
    found = 0
    
    with ThreadPoolExecutor(max_workers=WORKERS) as ex:
        results = ex.map(scan_pdf, files_to_scan)
        for result in results:
            scanned += 1
            if scanned % 500 == 0:
                elapsed = time.time() - start
                rate = scanned / elapsed
                eta = (len(files_to_scan) - scanned) / rate
                print(f'  [{scanned}/{len(files_to_scan)}] {rate:.1f} PDFs/s, ETA {eta/60:.0f}min, matches={found}')
            
            for phone in result['phones']:
                if phone in unknown_phones and phone not in matches:
                    matches[phone] = {
                        'filename': result['filename'],
                        'name_from_filename': result['name_from_filename'],
                    }
                    found += 1
                    if found <= 20:
                        print(f'  MATCH: {phone} -> {result["filename"][:60]} (name: {result["name_from_filename"]})')
            
            if found >= len(unknown_phones):
                print(f'  All unknowns matched! Stopping early.')
                break
    
    elapsed = time.time() - start
    print(f'\nDone: scanned {scanned}/{len(files_to_scan)} PDFs in {elapsed/60:.1f} minutes')
    print(f'Matched: {found}/{len(unknown_phones)} unknown phones')
    
    unmatched = unknown_phones - set(matches.keys())
    if unmatched:
        print(f'Unmatched phones ({len(unmatched)}):')
        for p in sorted(unmatched)[:20]:
            print(f'  {p}')
    
    with open(OUTPUT, 'w', encoding='utf-8') as f:
        json.dump(matches, f, ensure_ascii=False, indent=2)
    print(f'Results saved to {OUTPUT}')

if __name__ == '__main__':
    main()

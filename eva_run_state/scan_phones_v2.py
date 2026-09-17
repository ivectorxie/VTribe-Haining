import os, re, json, time, shutil, sys
from concurrent.futures import ThreadPoolExecutor
import pymupdf
import openpyxl

B_TABLE = r'S:\2 人事\4. 简历库\人才登记表20260916-B.xlsx'
SORTED_DIR = r'S:\2 人事\4. 简历库\Sorted'
UNSORTED_DIR = r'S:\2 人事\4. 简历库\Unsorted'
LOCAL_DIR = r'C:\Users\vecto\Documents\4.VTribe\_pdf_cache'
OUTPUT = r'C:\Users\vecto\Documents\4.VTribe\eva_run_state\phone_matches.json'

WORKERS = 20
CHINA_PHONE = re.compile(r'(?:\(\+86\))?\s*1[3-9]\d{9}')
INTL_PHONE = re.compile(r'\(\+\d+\)[\d\s\-]{6,15}')

PATTERNS = [
    re.compile(r'【(?P<title>[^_【】]+?)\s*[_\s](?P<city>[^_【】]*?)\s+(?P<salary>[^】]*?)】\s*(?P<name>[^】\d]+?)\s+(?P<exp>\d+年(?:以内|应届生|以上|毕业)?)(?:\s*-\s*Copy)?\.pdf', re.UNICODE),
    re.compile(r'【(?P<title>[^_【】]+?)\s*[_\s](?P<city>[^_【】]*?)】\s*(?P<name>[^】\d]+?)\s+(?P<exp>\d+年(?:以内|应届生|以上|毕业)?)(?:\s*-\s*Copy)?\.pdf', re.UNICODE),
    re.compile(r'(?P<name>[^_]+)_(?P<age>\d+岁|[\d\.]+)_(?P<title>[^_]+)_(?P<city>[^_]+)_.*\.pdf', re.UNICODE),
    re.compile(r'】\s*(?P<name>[^\d】\n]+?)\s+(?:\d+年|应届|\d+岁)(?:.*?)(?:\s*-\s*Copy)?\.pdf', re.UNICODE),
]

BAD_NAMES = {'未知', 'N/A', '', 'None', '费女士', '先生', '女士', '糖 糖'}

def extract_name_from_filename(filename):
    for pat in PATTERNS:
        m = pat.search(filename)
        if m:
            name = m.group('name').strip().strip('_- ')
            if name and name not in BAD_NAMES:
                return name
    return None

def normalize_phone(raw):
    digits = re.sub(r'\D', '', raw)
    if len(digits) == 11 and digits.startswith('1'):
        return f'(+86){digits}'
    if len(digits) > 11:
        country = raw[:raw.index('1')].strip() if '1' in raw else ''
        return f'{country}{digits[-11:]}'
    return None

def load_unknown_phones():
    wb = openpyxl.load_workbook(B_TABLE, read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    name_idx = headers.index('候选人')
    phone_idx = headers.index('电话1')
    unknowns = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        name = str(row[name_idx] or '').strip()
        if name in ('未知', 'N/A', '', 'None'):
            phone = str(row[phone_idx] or '').strip()
            if phone and phone not in ('N/A', '', 'None'):
                unknowns.add(phone)
    wb.close()
    print(f'Loaded {len(unknowns)} unknown phone numbers', flush=True)
    return unknowns

def copy_file(args):
    src, dst = args
    try:
        with open(src, 'rb') as f:
            data = f.read()
        with open(dst, 'wb') as f:
            f.write(data)
        return True
    except Exception as e:
        return False

def scan_pdf(filepath):
    filename = os.path.basename(filepath)
    try:
        doc = pymupdf.open(filepath)
        text = ''
        for page in doc:
            text += page.get_text()
        doc.close()
        phones = set()
        for m in CHINA_PHONE.finditer(text):
            norm = normalize_phone(m.group(0))
            if norm:
                phones.add(norm)
        for m in INTL_PHONE.finditer(text):
            raw = m.group(0).strip()
            if raw:
                phones.add(raw)
        name = extract_name_from_filename(filename)
        return {'filename': filename, 'phones': list(phones), 'name': name}
    except Exception as e:
        return {'filename': filename, 'phones': [], 'name': None, 'error': str(e)[:80]}

def main():
    start = time.time()
    unknown_phones = load_unknown_phones()

    os.makedirs(LOCAL_DIR, exist_ok=True)

    file_pairs = []
    for d, sub in [(SORTED_DIR, 'Sorted'), (UNSORTED_DIR, 'Unsorted')]:
        if not os.path.isdir(d):
            continue
        local_sub = os.path.join(LOCAL_DIR, sub)
        os.makedirs(local_sub, exist_ok=True)
        for f in os.listdir(d):
            if f.lower().endswith('.pdf'):
                src = os.path.join(d, f)
                dst = os.path.join(local_sub, f)
                file_pairs.append((src, dst, f, sub))

    print(f'Total PDFs to copy: {len(file_pairs)}', flush=True)

    to_copy = [(s, d) for s, d, f, sub in file_pairs if not os.path.exists(d)]
    print(f'Need to copy: {len(to_copy)} (already cached: {len(file_pairs) - len(to_copy)})', flush=True)

    if to_copy:
        print(f'Copying with {WORKERS} threads...', flush=True)
        copied = 0
        errors = 0
        with ThreadPoolExecutor(max_workers=WORKERS) as ex:
            results = list(ex.map(copy_file, to_copy))
        copied = sum(1 for r in results if r)
        errors = len(results) - copied
        elapsed_copy = time.time() - start
        print(f'Copied {copied} files ({errors} errors) in {elapsed_copy/60:.1f} min', flush=True)

    scan_start = time.time()
    local_files = [(os.path.join(LOCAL_DIR, sub, f), f) for _, _, f, sub in file_pairs]
    print(f'Scanning {len(local_files)} PDFs from local SSD...', flush=True)

    matches = {}
    scanned = 0

    with ThreadPoolExecutor(max_workers=WORKERS) as ex:
        results = ex.map(scan_pdf, [fp for fp, fn in local_files])
        for result in results:
            scanned += 1
            if scanned % 1000 == 0:
                elapsed = time.time() - scan_start
                rate = scanned / elapsed
                eta = (len(local_files) - scanned) / rate
                print(f'  [{scanned}/{len(local_files)}] {rate:.1f} PDFs/s, ETA {eta/60:.0f}min, matches={len(matches)}', flush=True)

            for phone in result['phones']:
                if phone in unknown_phones and phone not in matches:
                    matches[phone] = {
                        'filename': result['filename'],
                        'name': result['name'],
                    }
                    if len(matches) <= 30:
                        print(f'  MATCH #{len(matches)}: {phone} -> {result["filename"][:60]} (name: {result["name"]})', flush=True)

    elapsed_scan = time.time() - scan_start
    total_elapsed = time.time() - start

    print(f'\n=== RESULTS ===', flush=True)
    print(f'Scan: {scanned} PDFs in {elapsed_scan/60:.1f} min ({scanned/elapsed_scan:.1f} PDFs/s)', flush=True)
    print(f'Total: {total_elapsed/60:.1f} min', flush=True)
    print(f'Matched: {len(matches)}/{len(unknown_phones)} unknown phones', flush=True)

    unmatched = unknown_phones - set(matches.keys())
    print(f'Unmatched: {len(unmatched)}', flush=True)
    if unmatched:
        for p in sorted(unmatched)[:10]:
            print(f'  {p}', flush=True)

    with open(OUTPUT, 'w', encoding='utf-8') as f:
        json.dump(matches, f, ensure_ascii=False, indent=2)
    print(f'Saved to {OUTPUT}', flush=True)

if __name__ == '__main__':
    main()
